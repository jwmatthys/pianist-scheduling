import unittest
from io import BytesIO
from datetime import date, time
import tempfile
from uuid import uuid4
from unittest.mock import patch

from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import models, module_models, schemas
from app.database import Base, migrate_database
from app.routers import pianists
from app.services import availability_importer
from app.services.accompanist_availability import windows_to_accompanist_slots
from app.services.availability import (
    AvailabilityWindow,
    normalize_day,
    normalize_status,
    normalize_windows,
    parse_time_minutes,
)
from app.services.module_lifecycle import active_session_uuid, ensure_pianist_identity, module_revision, synchronize_lesson_identity


class AvailabilityValueTests(unittest.TestCase):
    def test_day_names_and_unambiguous_abbreviations_normalize(self):
        self.assertEqual(normalize_day("mOn."), "Monday")
        self.assertEqual(normalize_day("Thurs"), "Thursday")
        self.assertIsNone(normalize_day("T"))
        self.assertIsNone(normalize_day("Funday"))

    def test_time_parser_accepts_excel_times_and_common_text(self):
        self.assertEqual(parse_time_minutes(time(8, 0)), 480)
        self.assertEqual(parse_time_minutes(0.5), 720)
        self.assertEqual(parse_time_minutes("8:00 AM"), 480)
        self.assertEqual(parse_time_minutes("08:00"), 480)
        self.assertEqual(parse_time_minutes("1:30 PM"), 810)
        self.assertEqual(parse_time_minutes("24:00", allow_end_of_day=True), 1440)
        self.assertIsNone(parse_time_minutes("13:00 PM"))
        self.assertIsNone(parse_time_minutes("25:00"))
        self.assertIsNone(parse_time_minutes("8"))

    def test_status_vocabulary_is_case_insensitive_without_scoring_semantics(self):
        self.assertEqual(normalize_status(" tentative "), "Tentative")
        self.assertEqual(normalize_status("not sure"), None)

    def test_same_status_overlap_and_adjacency_merge(self):
        windows, issues = normalize_windows([
            AvailabilityWindow("Monday", 8 * 60, 11 * 60, "Available"),
            AvailabilityWindow("Monday", 10 * 60 + 30, 12 * 60, "Available"),
            AvailabilityWindow("Monday", 12 * 60, 13 * 60, "Available"),
        ])

        self.assertEqual([(item.start_minute, item.end_minute) for item in windows], [(480, 780)])
        self.assertEqual(issues, [])

    def test_conflicting_overlap_is_reported_and_not_merged(self):
        windows, issues = normalize_windows([
            AvailabilityWindow("Monday", 9 * 60, 12 * 60, "Available"),
            AvailabilityWindow("Monday", 10 * 60, 11 * 60, "Tentative", source_row=4),
        ])

        self.assertEqual(len(windows), 2)
        self.assertEqual([(issue.severity, issue.code, issue.row_number) for issue in issues], [
            ("error", "CONFLICTING_OVERLAP", 4),
        ])

    def test_duplicate_windows_are_deduplicated_with_warning(self):
        windows, issues = normalize_windows([
            AvailabilityWindow("Friday", 540, 600, "Unavailable", source_row=2),
            AvailabilityWindow("Friday", 540, 600, "Unavailable", source_row=3),
        ])

        self.assertEqual(len(windows), 1)
        self.assertEqual([(issue.severity, issue.code) for issue in issues], [("warning", "DUPLICATE_WINDOW")])

    def test_accompanist_adapter_preserves_existing_half_hour_fit_and_rejects_rounding(self):
        windows = [AvailabilityWindow("Monday", 8 * 60, 10 * 60, "Available")]
        slots, issues = windows_to_accompanist_slots(windows)
        availability = {windows[0].day: {slot.slot_start_minute: slot.status for slot in slots}}

        from app.services.scheduling import FIT_FULL, get_fit

        self.assertEqual(issues, [])
        self.assertEqual(get_fit("Monday", 8 * 60, 9 * 60, availability)[0], FIT_FULL)

        rounded_slots, invalid = windows_to_accompanist_slots([
            AvailabilityWindow("Monday", 8 * 60 + 15, 10 * 60, "Available", source_row=7),
        ])
        self.assertEqual(rounded_slots, [])
        self.assertEqual([(issue.code, issue.row_number) for issue in invalid], [
            ("UNSUPPORTED_ACCOMPANIST_BOUNDARY", 7),
        ])
        unavailable_slots, invalid = windows_to_accompanist_slots([
            AvailabilityWindow("Monday", 8 * 60, 10 * 60, "Unavailable"),
        ])
        self.assertEqual(unavailable_slots, [])
        self.assertEqual(invalid, [])


class AvailabilityImportServiceTests(unittest.TestCase):
    def setUp(self):
        availability_importer._STAGED_UPLOADS.clear()
        availability_importer._STAGED_PREVIEWS.clear()
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        migrate_database(self.engine)
        self.db = Session(self.engine)
        self.db.add_all([
            models.Pianist(name="Ari Example"),
            models.Pianist(name="Bea Sample"),
        ])
        self.db.commit()

    def tearDown(self):
        availability_importer._STAGED_UPLOADS.clear()
        availability_importer._STAGED_PREVIEWS.clear()
        self.db.close()
        self.engine.dispose()

    def pianist_id(self, name):
        return self.db.query(models.Pianist).filter_by(name=name).one().id

    def csv_upload(self, content):
        return availability_importer.inspect_upload("synthetic.csv", content.encode())

    def normalized_mapping(self, inspection, person="Pianist Name", email=None):
        return schemas.AvailabilityImportMapping(
            layout="normalized",
            person_name_column=person,
            email_column=email,
            day_column="Weekday",
            start_column="Start",
            end_column="End",
            status_column="Status",
        )

    def preview_csv(self, content, *, person="Pianist Name", email=None):
        inspection = self.csv_upload(content)
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=inspection.upload_token,
                mapping=self.normalized_mapping(inspection, person, email),
            ),
            self.db,
        )
        return inspection, preview

    def add_jury_availability(self, person_uuid, jury_date):
        session_uuid = active_session_uuid(self.db)
        self.db.add(module_models.JuryPianistAvailabilityDeclaration(
            session_uuid=session_uuid,
            pianist_person_uuid=person_uuid,
            jury_date=jury_date,
            is_complete=True,
        ))
        self.db.add(module_models.JuryPianistAvailableWindow(
            session_uuid=session_uuid,
            pianist_person_uuid=person_uuid,
            jury_date=jury_date,
            start_minute=480,
            end_minute=540,
        ))

    def test_import_groups_repeated_normalized_names_and_replaces_old_roster(self):
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "  Ari   Example  ,Mon,08:00,09:30,Available\n"
            "Ari Example,Monday,14:00,15:00,Tentative\n"
        )

        self.assertTrue(preview.can_apply)
        self.assertEqual(preview.rows_processed, 2)
        self.assertEqual(preview.existing_pianist_count, 2)
        self.assertEqual(preview.incoming_pianist_count, 1)
        self.assertEqual(preview.valid_window_count, 2)
        self.assertIn("FULL_ROSTER_REPLACEMENT", {warning.code for warning in preview.warnings})
        self.assertEqual([person.pianist_name for person in preview.pianists], ["Ari Example"])
        self.assertEqual([window.status for window in preview.windows], ["Available", "Tentative"])
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 0)
        self.assertNotIn("pianist_code", preview.model_dump())
        self.assertNotIn("pianist_id", preview.model_dump())

        result = availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )

        self.assertEqual((result.pianists_removed, result.pianists_created), (2, 1))
        self.assertEqual(result.slots_created, 5)
        self.assertEqual(self.db.query(models.Pianist).one().name, "Ari Example")
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 5)

    def test_import_without_identifier_column_creates_fresh_people_with_defaults(self):
        self.db.query(models.Pianist).delete()
        self.db.commit()
        rows = [f"Person {index},,Mon,08:00,09:00,Available" for index in range(1, 7)]
        inspection, preview = self.preview_csv(
            "Pianist Name,Email,Weekday,Start,End,Status\n" + "\n".join(rows) + "\n",
        )

        self.assertTrue(preview.can_apply)
        self.assertEqual(preview.incoming_pianist_count, 6)
        self.assertNotIn("person_id_column", inspection.suggested_normalized)
        self.assertTrue(all(item.action == "new" and item.email == "" for item in preview.pianists))
        self.assertTrue(all(item.max_hours_per_week == 40 for item in preview.pianists))

        result = availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )

        pianists = self.db.query(models.Pianist).order_by(models.Pianist.name).all()
        mappings = self.db.query(module_models.AccompanistPianistIdentity).all()
        self.assertEqual((result.pianists_created, result.pianists_removed), (6, 0))
        self.assertEqual(len(pianists), 6)
        self.assertEqual(len(mappings), 6)
        self.assertEqual(len({mapping.person_uuid for mapping in mappings}), 6)
        self.assertTrue(all(pianist.email == "" and pianist.max_hours_per_week == 40 for pianist in pianists))
        self.assertTrue(all(pianist.pianist_code is None for pianist in pianists))
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 12)
        self.assertTrue(all(pianist.availability_complete for pianist in pianists))

    def test_template_and_user_models_have_no_pianist_identifier(self):
        from pathlib import Path

        template = Path(__file__).resolve().parents[2] / "frontend" / "public" / "availability-template.csv"
        self.assertEqual(template.read_text(encoding="utf-8").splitlines()[0], "Pianist Name,Day,Start,End,Status")
        self.assertNotIn("pianist_code", schemas.PianistCreate.model_fields)
        self.assertNotIn("pianist_code", schemas.PianistUpdate.model_fields)
        self.assertNotIn("pianist_code", schemas.PianistOut.model_fields)
        self.assertNotIn("person_id_column", schemas.AvailabilityImportMapping.model_fields)
        with self.assertRaises(Exception):
            schemas.AvailabilityImportMapping(
                layout="normalized",
                person_name_column="Pianist Name",
                person_id_column="Pianist ID",
            )

    def test_imported_defaults_remain_editable_through_human_facing_api(self):
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\nEditable Import,Mon,08:00,09:00,Available\n"
        )
        availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )
        imported = self.db.query(models.Pianist).filter_by(name="Editable Import").one()
        self.assertEqual((imported.email, imported.max_hours_per_week), ("", 40))
        updated = pianists.update_pianist(imported.id, schemas.PianistUpdate(
            name="Edited Import",
            email="edited@example.invalid",
            max_hours_per_week=12,
        ), self.db)
        self.assertEqual(
            (updated.name, updated.email, updated.max_hours_per_week),
            ("Edited Import", "edited@example.invalid", 12),
        )

    def test_manual_pianist_creation_generates_internal_identity_without_id_input(self):
        created = pianists.create_pianist(schemas.PianistCreate(name="Manual Pianist"), self.db)
        self.assertIsNone(created.pianist_code)
        self.assertNotIn("pianist_code", schemas.PianistOut.model_validate(created).model_dump())
        mapping = self.db.get(module_models.AccompanistPianistIdentity, created.id)
        self.assertIsNotNone(mapping)
        self.assertIsNotNone(self.db.get(module_models.PersonIdentity, mapping.person_uuid))
        with self.assertRaises(Exception):
            schemas.PianistCreate(name="Invalid Input", pianist_code="P001")

    def test_invalid_window_blocks_new_pianist_creation_atomically(self):
        ari_id = self.pianist_id("Ari Example")
        ari_uuid = ensure_pianist_identity(self.db, self.db.get(models.Pianist, ari_id))
        self.add_jury_availability(ari_uuid, date(2027, 12, 14))
        self.db.commit()
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "Would Be New,Mon,08:00,09:00,Available\n"
            "Would Be New,Tue,10:00,,Available\n"
        )
        self.assertFalse(preview.can_apply)
        with self.assertRaises(availability_importer.AvailabilityImportError):
            availability_importer.apply_import(
                schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
                self.db,
            )
        self.assertEqual(self.db.query(models.Pianist).filter_by(name="Would Be New").count(), 0)
        self.assertEqual(self.db.query(module_models.AccompanistPianistIdentity).count(), 1)
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 0)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailabilityDeclaration).count(), 1)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailableWindow).count(), 1)

    def test_cancelled_replacement_preserves_jury_availability(self):
        ari_id = self.pianist_id("Ari Example")
        ari_uuid = ensure_pianist_identity(self.db, self.db.get(models.Pianist, ari_id))
        self.add_jury_availability(ari_uuid, date(2027, 12, 14))
        self.db.commit()
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\nNew Pianist,Mon,08:00,09:00,Available\n"
        )

        with self.assertRaises(availability_importer.AvailabilityImportError) as raised:
            availability_importer.apply_import(
                schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=False),
                self.db,
            )
        self.assertEqual(raised.exception.code, "CONFIRMATION_REQUIRED")
        self.assertEqual(self.db.query(models.Pianist).count(), 2)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailabilityDeclaration).count(), 1)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailableWindow).count(), 1)

    def test_failed_replacement_rolls_back_jury_cleanup_and_accompanist_changes(self):
        ari_id = self.pianist_id("Ari Example")
        ari_uuid = ensure_pianist_identity(self.db, self.db.get(models.Pianist, ari_id))
        self.add_jury_availability(ari_uuid, date(2027, 12, 14))
        jury_configuration = self.db.get(module_models.JuryConfiguration, active_session_uuid(self.db))
        initial_jury_revision = jury_configuration.input_revision
        lesson = models.Lesson(
            student="Rollback Student",
            day="Monday",
            start_minute=480,
            end_minute=510,
            assigned_pianist_id=ari_id,
        )
        self.db.add_all([
            lesson,
            models.AvailabilitySlot(
                pianist_id=ari_id,
                day="Monday",
                slot_start_minute=480,
                status="Available",
            ),
        ])
        self.db.commit()
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\nNew Pianist,Mon,08:00,09:00,Available\n"
        )

        with patch.object(availability_importer, "bump_accompanist_revision", side_effect=RuntimeError("forced failure")):
            with self.assertRaises(availability_importer.AvailabilityImportError) as raised:
                availability_importer.apply_import(
                    schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
                    self.db,
                )

        self.assertEqual(raised.exception.code, "APPLY_FAILED")
        self.db.expire_all()
        self.assertEqual(self.db.query(models.Pianist).count(), 2)
        self.assertEqual(self.db.query(models.Pianist).filter_by(name="New Pianist").count(), 0)
        self.assertEqual(self.db.get(models.Lesson, lesson.id).assigned_pianist_id, ari_id)
        self.assertEqual(self.db.query(models.AvailabilitySlot).filter_by(pianist_id=ari_id).count(), 1)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailabilityDeclaration).count(), 1)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailableWindow).count(), 1)
        self.assertEqual(
            self.db.get(module_models.JuryConfiguration, active_session_uuid(self.db)).input_revision,
            initial_jury_revision,
        )

    def test_reimport_same_name_replaces_identity_instead_of_matching_old_roster(self):
        csv = "Pianist Name,Weekday,Start,End,Status\n"
        first_inspection, first_preview = self.preview_csv(
            csv + "Repeat Import,Mon,08:00,09:00,Available\n",
        )
        availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=first_preview.preview_token, confirmed=True),
            self.db,
        )
        first_pianist = self.db.query(models.Pianist).filter_by(name="Repeat Import").one()
        first_identity = self.db.get(module_models.AccompanistPianistIdentity, first_pianist.id).person_uuid
        second_preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=first_inspection.upload_token,
                mapping=self.normalized_mapping(first_inspection),
            ),
            self.db,
        )
        self.assertEqual((second_preview.existing_pianist_count, second_preview.incoming_pianist_count), (1, 1))
        availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=second_preview.preview_token, confirmed=True),
            self.db,
        )
        second_pianist = self.db.query(models.Pianist).filter_by(name="Repeat Import").one()
        second_identity = self.db.get(module_models.AccompanistPianistIdentity, second_pianist.id).person_uuid
        self.assertNotEqual(first_identity, second_identity)
        self.assertEqual(self.db.query(models.Pianist).count(), 1)

    def test_xlsx_actual_time_cells_multiple_sheets_and_forms_wide_suggestions(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Form Responses"
        headers = ["Full Name"]
        row = ["Ari Example"]
        for day in models.DAYS_ORDER:
            if day == "Monday":
                headers.extend(["Monday Start Time", "Monday End Time", "Monday Start 2", "Monday End 2"])
                row.extend([time(8, 0), time(9, 30), time(14, 0), time(15, 0)])
            else:
                headers.extend([f"{day} Start Time", f"{day} End Time"])
                row.extend([None, None])
        sheet.append(headers)
        sheet.append(row)
        workbook.create_sheet("Unused Sheet").append(["Other", "Data"])
        contents = BytesIO()
        workbook.save(contents)

        inspection = availability_importer.inspect_upload("forms.xlsx", contents.getvalue())

        self.assertEqual(inspection.sheets, ["Form Responses", "Unused Sheet"])
        self.assertEqual(len(inspection.suggested_wide_windows["Monday"]), 2)
        other_sheet = availability_importer.inspect_sheet(inspection.upload_token, "Unused Sheet")
        self.assertEqual(other_sheet.columns, ["Other", "Data"])
        mapping = schemas.AvailabilityImportMapping(
            layout="wide",
            person_name_column="Full Name",
            wide_status="Available",
            wide_windows=inspection.suggested_wide_windows,
        )
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(upload_token=inspection.upload_token, mapping=mapping),
            self.db,
        )

        self.assertEqual(preview.valid_window_count, 2)
        self.assertEqual([window.start_minute for window in preview.windows], [480, 840])
        self.assertTrue(preview.can_apply)

    def test_normalized_xlsx_reads_excel_time_cells(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Pianist Name", "Weekday", "Start", "End", "Status"])
        sheet.append(["Ari Example", "Mon", time(8, 0), time(9, 30), "Available"])
        contents = BytesIO()
        workbook.save(contents)
        inspection = availability_importer.inspect_upload("normalized.xlsx", contents.getvalue())
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=inspection.upload_token,
                mapping=schemas.AvailabilityImportMapping(
                    layout="normalized",
                    person_name_column="Pianist Name",
                    day_column="Weekday",
                    start_column="Start",
                    end_column="End",
                    status_column="Status",
                ),
            ),
            self.db,
        )

        self.assertEqual(preview.valid_window_count, 1)
        self.assertEqual((preview.windows[0].start_minute, preview.windows[0].end_minute), (480, 570))

    def test_meaningfully_different_names_create_separate_pianists(self):
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "Alex Smith,Mon,08:00,09:00,Available\n"
            "Alex J. Smith,Mon,09:00,10:00,Available\n"
        )

        self.assertTrue(preview.can_apply)
        self.assertEqual(preview.incoming_pianist_count, 2)
        self.assertEqual({person.pianist_name for person in preview.pianists}, {"Alex Smith", "Alex J. Smith"})
        availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )
        self.assertEqual(self.db.query(models.Pianist).count(), 2)

    def test_invalid_mapping_day_time_range_status_and_conflict_block_all_writes(self):
        missing_mapping_file = self.csv_upload(
            "Pianist Name,Weekday,Start,End,Status\nAri Example,Mon,08:00,09:00,Available\n"
        )
        missing_mapping = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=missing_mapping_file.upload_token,
                mapping=schemas.AvailabilityImportMapping(layout="normalized"),
            ),
            self.db,
        )
        self.assertIn("MISSING_MAPPING", {issue.code for issue in missing_mapping.errors})
        self.assertFalse(missing_mapping.can_apply)

        _, missing_mapping = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\nAri Example,Mon,08:00,09:00,Available\n",
            person="Does Not Exist",
        )
        self.assertIn("INVALID_MAPPING", {issue.code for issue in missing_mapping.errors})

        _, invalid_rows = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "Ari Example,Funday,08:00,09:00,Available\n"
            "Ari Example,Mon,25:00,26:00,Available\n"
            "Ari Example,Tue,10:00,09:00,Available\n"
            "Ari Example,Wed,08:00,09:00,Maybe\n"
        )
        self.assertEqual(
            {issue.code for issue in invalid_rows.errors},
            {"INVALID_DAY", "INVALID_START_TIME", "INVALID_END_TIME", "INVALID_TIME_RANGE", "INVALID_STATUS"},
        )

        _, conflict = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "Ari Example,Mon,09:00,12:00,Available\n"
            "Ari Example,Mon,10:00,11:00,Tentative\n"
        )
        self.assertIn("CONFLICTING_OVERLAP", {issue.code for issue in conflict.errors})
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 0)

    def test_full_replacement_clears_old_assignments_availability_and_absent_pianists(self):
        ari_id = self.pianist_id("Ari Example")
        bea_id = self.pianist_id("Bea Sample")
        old_person_uuid = ensure_pianist_identity(self.db, self.db.get(models.Pianist, ari_id))
        bea_person_uuid = ensure_pianist_identity(self.db, self.db.get(models.Pianist, bea_id))
        first_date = date(2027, 12, 14)
        second_date = date(2027, 12, 15)
        lesson = models.Lesson(
            student="Synthetic Student",
            day="Monday",
            start_minute=480,
            end_minute=510,
            assigned_pianist_id=ari_id,
        )
        self.db.add(lesson)
        self.db.flush()
        student_person_uuid = synchronize_lesson_identity(self.db, lesson)
        lesson_identity = self.db.get(module_models.AccompanistLessonIdentity, lesson.id)
        jury_panel = module_models.JuryPanel(
            session_uuid=active_session_uuid(self.db),
            panel_name="Panel A",
            earliest_start_minute=480,
            preferred_start_minute=540,
            jury_length_minutes=10,
        )
        second_panel = module_models.JuryPanel(
            session_uuid=active_session_uuid(self.db),
            panel_name="Panel B",
            earliest_start_minute=480,
            preferred_start_minute=540,
            jury_length_minutes=10,
        )
        self.db.add_all([jury_panel, second_panel])
        self.db.flush()
        self.db.add_all([
            module_models.JuryPanelDate(
                panel_uuid=jury_panel.panel_uuid,
                session_uuid=active_session_uuid(self.db),
                jury_date=first_date,
            ),
            module_models.JuryPanelDate(
                panel_uuid=second_panel.panel_uuid,
                session_uuid=active_session_uuid(self.db),
                jury_date=second_date,
            ),
        ])
        self.db.add_all([
            module_models.AccompanistLessonJuryRequirement(lesson_id=lesson.id, jury_required=True),
            module_models.JuryLessonEntry(
                session_uuid=active_session_uuid(self.db),
                source_lesson_uuid=lesson_identity.lesson_uuid,
                student_person_uuid=student_person_uuid,
                jury_required=True,
                panel_uuid=jury_panel.panel_uuid,
            ),
        ])
        revision = module_revision(self.db)
        initial_source_revision = revision.source_revision
        previous_result_uuid = str(uuid4())
        self.db.add(module_models.ModuleResult(
            result_uuid=previous_result_uuid,
            session_uuid=revision.session_uuid,
            module_id="accompanists",
            contract_id="accompanist.assignment-result",
            contract_version=2,
            payload_schema_version=1,
            result_version=1,
            source_revision=revision.source_revision,
            state="finalized",
            payload_json="{}",
            payload_sha256="0" * 64,
        ))
        revision.current_result_uuid = previous_result_uuid
        for pianist_uuid in (old_person_uuid, bea_person_uuid):
            for jury_date in (first_date, second_date):
                self.add_jury_availability(pianist_uuid, jury_date)
        self.db.add_all([
            models.AvailabilitySlot(pianist_id=ari_id, day="Monday", slot_start_minute=480, status="Available"),
            models.AvailabilitySlot(pianist_id=ari_id, day="Monday", slot_start_minute=510, status="Tentative"),
            models.AvailabilitySlot(pianist_id=ari_id, day="Tuesday", slot_start_minute=540, status="Available"),
            models.AvailabilitySlot(pianist_id=bea_id, day="Monday", slot_start_minute=600, status="Available"),
        ])
        self.db.commit()
        self.assertEqual(self.db.query(module_models.JuryPianistAvailabilityDeclaration).count(), 4)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailableWindow).count(), 4)
        jury_configuration = self.db.get(module_models.JuryConfiguration, active_session_uuid(self.db))
        initial_jury_revision = jury_configuration.input_revision
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\n"
            "Ari Example,Tue,10:00,11:00,Available\n"
            "Ari Example,Thu,09:00,10:00,Tentative\n"
        )

        self.assertEqual(preview.existing_pianist_count, 2)
        self.assertEqual(preview.existing_assignment_count, 1)
        self.assertEqual(preview.existing_slots_in_scope, 4)
        result = availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )

        imported = self.db.query(models.Pianist).one()
        new_person_uuid = self.db.get(
            module_models.AccompanistPianistIdentity,
            imported.id,
        ).person_uuid
        slots = self.db.query(models.AvailabilitySlot).order_by(
            models.AvailabilitySlot.day, models.AvailabilitySlot.slot_start_minute
        ).all()
        self.assertEqual(result.pianists_removed, 2)
        self.assertEqual(result.assignments_cleared, 1)
        self.assertEqual(result.slots_replaced, 4)
        self.assertEqual(result.slots_created, 4)
        self.assertEqual(result.jury_availability_windows_removed, 4)
        self.assertEqual(
            self.db.get(module_models.JuryConfiguration, active_session_uuid(self.db)).input_revision,
            initial_jury_revision + 1,
        )
        self.assertEqual(module_revision(self.db).source_revision, initial_source_revision + 1)
        self.assertIsNone(module_revision(self.db).current_result_uuid)
        self.assertEqual(self.db.get(module_models.ModuleResult, previous_result_uuid).state, "superseded")
        self.assertIsNone(self.db.get(models.Lesson, lesson.id).assigned_pianist_id)
        self.assertNotEqual(new_person_uuid, old_person_uuid)
        self.assertEqual(self.db.query(models.Pianist).filter_by(name="Bea Sample").count(), 0)
        self.assertEqual(
            self.db.get(module_models.AccompanistLessonIdentity, lesson.id).lesson_uuid,
            lesson_identity.lesson_uuid,
        )
        self.assertEqual(
            self.db.get(
                module_models.JuryLessonEntry,
                (active_session_uuid(self.db), lesson_identity.lesson_uuid),
            ).panel_uuid,
            jury_panel.panel_uuid,
        )
        self.assertTrue(self.db.get(module_models.AccompanistLessonJuryRequirement, lesson.id).jury_required)
        self.assertEqual(self.db.query(models.Lesson).count(), 1)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailabilityDeclaration).count(), 0)
        self.assertEqual(self.db.query(module_models.JuryPianistAvailableWindow).count(), 0)
        self.assertEqual(self.db.query(module_models.JuryPanel).count(), 2)
        self.assertEqual(self.db.query(module_models.JuryPanelDate).count(), 2)
        self.assertEqual(
            {row.jury_date for row in self.db.query(module_models.JuryPanelDate).all()},
            {first_date, second_date},
        )
        self.assertEqual(
            [(slot.pianist_id, slot.day, slot.slot_start_minute, slot.status) for slot in slots],
            [
                (imported.id, "Thursday", 540, "Tentative"),
                (imported.id, "Thursday", 570, "Tentative"),
                (imported.id, "Tuesday", 600, "Available"),
                (imported.id, "Tuesday", 630, "Available"),
            ],
        )
        self.assertTrue(imported.availability_complete)
        self.assertNotIn("Monday", {slot.day for slot in slots})

        from app.services.scheduling import FIT_FULL, FIT_NONE, get_fit

        availability_by_day = {}
        for slot in slots:
            availability_by_day.setdefault(slot.day, {})[slot.slot_start_minute] = slot.status
        self.assertEqual(get_fit("Monday", 480, 510, availability_by_day)[0], FIT_NONE)
        self.assertEqual(get_fit("Wednesday", 480, 510, availability_by_day)[0], FIT_NONE)

        pianists.set_availability(
            imported.id,
            schemas.AvailabilityBulkIn(slots=[
                schemas.AvailabilitySlotIn(day="Monday", slot_start_minute=480, status="Available"),
            ]),
            self.db,
        )
        availability = {"Monday": {slot.slot_start_minute: slot.status for slot in self.db.query(models.AvailabilitySlot).filter_by(pianist_id=imported.id).all()}}
        self.assertTrue(self.db.get(models.Pianist, imported.id).availability_complete)
        self.assertEqual(get_fit("Monday", 480, 510, availability)[0], FIT_FULL)

    def test_valid_zero_window_submission_clears_old_slots_and_marks_week_complete(self):
        ari_id = self.pianist_id("Ari Example")
        self.db.add(models.AvailabilitySlot(pianist_id=ari_id, day="Tuesday", slot_start_minute=480, status="Available"))
        self.db.commit()
        _, preview = self.preview_csv(
            "Pianist Name,Weekday,Start,End,Status\nAri Example,,,,\n",
        )
        self.assertTrue(preview.can_apply)
        self.assertEqual(preview.valid_window_count, 0)
        self.assertEqual(preview.errors, [])

        result = availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )
        slots = self.db.query(models.AvailabilitySlot).filter_by(pianist_id=ari_id).all()
        self.assertEqual(result.slots_created, 0)
        self.assertEqual(slots, [])
        self.assertTrue(self.db.get(models.Pianist, ari_id).availability_complete)

    def test_malformed_complete_submission_does_not_replace_or_complete_availability(self):
        ari_id = self.pianist_id("Ari Example")
        self.db.add(models.AvailabilitySlot(
            pianist_id=ari_id,
            day="Tuesday",
            slot_start_minute=600,
            status="Available",
        ))
        self.db.commit()
        inspection = self.csv_upload(
            "Pianist Name,Weekday,Start,End,Status\nAri Example,Tuesday,10:00,,Available\n"
        )
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=inspection.upload_token,
                mapping=self.normalized_mapping(inspection),
            ),
            self.db,
        )

        self.assertFalse(preview.can_apply)
        self.assertIn("INVALID_END_TIME", {issue.code for issue in preview.errors})
        with self.assertRaises(availability_importer.AvailabilityImportError):
            availability_importer.apply_import(
                schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
                self.db,
            )
        self.assertEqual(
            [(slot.day, slot.slot_start_minute, slot.status) for slot in self.db.query(models.AvailabilitySlot).all()],
            [("Tuesday", 600, "Available")],
        )
        self.assertFalse(self.db.get(models.Pianist, ari_id).availability_complete)

    def test_wide_import_missing_a_weekday_mapping_is_structurally_incomplete(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Name", "Monday Start", "Monday End"])
        sheet.append(["Ari Example", time(8, 0), time(9, 0)])
        content = BytesIO()
        workbook.save(content)
        inspection = availability_importer.inspect_upload("incomplete-wide.xlsx", content.getvalue())
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(
                upload_token=inspection.upload_token,
                mapping=schemas.AvailabilityImportMapping(
                    layout="wide",
                    person_name_column="Name",
                    wide_windows=inspection.suggested_wide_windows,
                ),
            ),
            self.db,
        )

        self.assertFalse(preview.can_apply)
        self.assertGreaterEqual(
            sum(issue.code == "MISSING_MAPPING" for issue in preview.errors),
            len(models.DAYS_ORDER) - 1,
        )
        self.assertEqual(self.db.query(models.AvailabilitySlot).count(), 0)

    def test_complete_wide_blank_days_are_unavailable_without_warnings(self):
        workbook = Workbook()
        sheet = workbook.active
        headers = ["Name"]
        row = ["Ari Example"]
        for day in models.DAYS_ORDER:
            headers.extend([f"{day} Start", f"{day} End"])
            if day == "Monday":
                row.extend([time(8, 0), time(9, 0)])
            else:
                row.extend([None, None])
        sheet.append(headers)
        sheet.append(row)
        content = BytesIO()
        workbook.save(content)
        inspection = availability_importer.inspect_upload("wide.xlsx", content.getvalue())
        mapping = schemas.AvailabilityImportMapping(
            layout="wide",
            person_name_column="Name",
            wide_windows=inspection.suggested_wide_windows,
        )
        preview = availability_importer.preview_import(
            schemas.AvailabilityImportPreviewRequest(upload_token=inspection.upload_token, mapping=mapping),
            self.db,
        )
        self.assertEqual(preview.pianists[0].days, ["Monday"])
        self.assertEqual(preview.valid_window_count, 1)
        self.assertFalse(any(issue.code in {"MISSING_VS_UNAVAILABLE", "MISSING_DAY_DATA", "NO_AVAILABILITY"} for issue in preview.warnings))
        self.assertTrue(preview.can_apply)

        availability_importer.apply_import(
            schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
            self.db,
        )
        slots = self.db.query(models.AvailabilitySlot).filter_by(pianist_id=self.pianist_id("Ari Example")).all()
        self.assertEqual(len(slots), 2)
        self.assertTrue(all(slot.status == "Available" for slot in slots))
        self.assertTrue(self.db.get(models.Pianist, self.pianist_id("Ari Example")).availability_complete)

    def test_imported_availability_survives_session_archive_round_trip(self):
        from app.services.session_files import create_new_session, export_session_archive, restore_session

        with tempfile.TemporaryDirectory() as directory:
            database_path = f"{directory}/session.sqlite3"
            engine = create_engine(f"sqlite:///{database_path}", connect_args={"check_same_thread": False})
            migrate_database(engine)
            session = Session(engine)
            inspection = availability_importer.inspect_upload(
                "round-trip.csv",
                b"Pianist Name,Weekday,Start,End,Status\nRound Trip Pianist,Mon,08:00,09:00,Available\n",
            )
            preview = availability_importer.preview_import(
                schemas.AvailabilityImportPreviewRequest(
                    upload_token=inspection.upload_token,
                    mapping=schemas.AvailabilityImportMapping(
                        layout="normalized",
                        person_name_column="Pianist Name",
                        day_column="Weekday",
                        start_column="Start",
                        end_column="End",
                        status_column="Status",
                    ),
                ),
                session,
            )
            availability_importer.apply_import(
                schemas.AvailabilityImportApplyRequest(preview_token=preview.preview_token, confirmed=True),
                session,
            )
            imported_pianist = session.query(models.Pianist).one()
            imported_person_uuid = session.get(
                module_models.AccompanistPianistIdentity,
                imported_pianist.id,
            ).person_uuid
            session.close()
            archive = export_session_archive(engine)
            create_new_session(
                engine,
                schemas.SessionMetadataIn(
                    institution_name="Synthetic University",
                    program_name="Music",
                    term_label="Spring 2027",
                ),
            )
            restore_session(engine, archive)
            with Session(engine) as restored_db:
                restored_slots = restored_db.query(models.AvailabilitySlot).all()
                self.assertEqual(
                    [(slot.day, slot.slot_start_minute, slot.status) for slot in restored_slots],
                    [("Monday", minute, "Available") for minute in (480, 510)],
                )
                restored_pianist = restored_db.query(models.Pianist).one()
                self.assertTrue(restored_pianist.availability_complete)
                identity_mapping = restored_db.get(module_models.AccompanistPianistIdentity, restored_pianist.id)
                self.assertIsNotNone(identity_mapping)
                self.assertEqual(identity_mapping.person_uuid, imported_person_uuid)
                self.assertIsNotNone(restored_db.get(module_models.PersonIdentity, identity_mapping.person_uuid))
            engine.dispose()


if __name__ == "__main__":
    unittest.main()