"""
Sample Workbook Generator

Writes three sample input workbooks matching the schema documented in the
README, so the scheduling scripts can be exercised without real data:

  - lesson_information.xlsx    ("Lessons" sheet)
  - pianist_availability.xlsx  ("Pianist - <Name>" sheets)
  - jury_information.xlsx      ("Jury Information" sheet)
"""

from openpyxl import Workbook
from datetime import time

DAYS_ORDER = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

STUDENTS = [
    ("Alice Chen",   "Piano",  "Monday",    time(9, 0),  time(9, 30),  "Room 101", "Area A"),
    ("Ben Torres",   "Violin", "Monday",    time(10, 0), time(10, 30), "Room 102", "Area A"),
    ("Cara Nguyen",  "Cello",  "Tuesday",   time(9, 0),  time(9, 45),  "Room 101", "Area B"),
    ("David Osei",   "Flute",  "Tuesday",   time(11, 0), time(11, 30), "Room 103", "Area B"),
    ("Ella Ramirez", "Voice",  "Wednesday", time(13, 0), time(13, 30), "Room 101", "Area A"),
    ("Finn Walsh",   "Piano",  "Thursday",  time(9, 0),  time(9, 30),  "Room 102", "Area B"),
]

PIANISTS = [
    ("Grace Kim",      12),
    ("Henry Ford",     10),
]


def write_lessons(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Lessons"
    headers = ["Lesson Teacher Name", "Student Name", "Lesson Day",
               "Lesson Start Time", "Lesson End Time", "Lesson Location",
               "Instrument", "Need Pianist", "Required Accompanist",
               "Jury", "Area"]
    ws.append(headers)
    for i, (student, instrument, day, start, end, room, area) in enumerate(STUDENTS):
        teacher = f"Teacher {i % 2 + 1}"
        ws.append([teacher, student, day, start, end, room, instrument,
                   1, "", 1, area])
    wb.save(path)


def write_pianists(path):
    wb = Workbook()
    wb.remove(wb.active)
    for name, max_hours in PIANISTS:
        ws = wb.create_sheet(f"Pianist - {name}")
        ws.append(["Name:", name])
        ws.append(["Max weekly hours:", max_hours])
        ws.append([])
        ws.append(["Time", *DAYS_ORDER, "Jury Day Availability"])
        t = time(8, 0)
        for slot in range(20):  # 8:00 - 18:00 in 30-minute slots
            minutes = t.hour * 60 + t.minute + slot * 30
            row_time = time(minutes // 60, minutes % 60)
            row = [row_time] + ["Available"] * len(DAYS_ORDER) + ["Available"]
            ws.append(row)
    wb.save(path)


def write_jury_info(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Jury Information"
    ws.append(["Area", "Start Time", "Jury Length", "Location",
               "Hourly Break", "Lunch Break"])
    ws.append(["Area A", time(9, 0), 15, "Room 101", 1, 1])
    ws.append(["Area B", time(9, 0), 20, "Room 102", 1, 1])
    wb.save(path)


if __name__ == "__main__":
    write_lessons("lesson_information.xlsx")
    write_pianists("pianist_availability.xlsx")
    write_jury_info("jury_information.xlsx")
    print("Wrote lesson_information.xlsx, pianist_availability.xlsx, jury_information.xlsx")
